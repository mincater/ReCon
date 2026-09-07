"""DOCX document loader using python-docx."""

import os
from typing import Any, Dict, List


def load(file_path: str) -> List[Dict[str, Any]]:
    """Loads a DOCX document and extracts paragraphs and tabular content.

    Args:
        file_path: Path to the DOCX file.

    Returns:
        List containing {"text": str, "source": str, "page": None}.

    Raises:
        FileNotFoundError: If the file does not exist.
        ImportError: If python-docx is not installed.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"DOCX file not found: {file_path}")

    try:
        import docx
    except ImportError as e:
        raise ImportError(
            "python-docx is required for loading DOCX documents. Install via requirements.txt."
        ) from e

    doc = docx.Document(file_path)
    content_blocks: List[str] = []

    # Extract paragraphs
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            content_blocks.append(text)

    # Extract tables
    for table in doc.tables:
        table_rows: List[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                table_rows.append(" | ".join(cells))
        if table_rows:
            content_blocks.append("\n".join(table_rows))

    full_text = "\n\n".join(content_blocks)
    return [{"text": full_text, "source": file_path, "page": None}]
